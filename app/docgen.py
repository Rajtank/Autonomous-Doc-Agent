"""Build a styled Word document from the agent's plan and sections."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .schemas import Plan, Section, TableData


def _add_table(doc: Document, table: TableData) -> None:
    if table.caption:
        caption = doc.add_paragraph()
        run = caption.add_run(table.caption)
        run.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    if not table.headers:
        return

    t = doc.add_table(rows=1 + len(table.rows), cols=len(table.headers))
    try:
        t.style = "Light Grid Accent 1"
    except KeyError:
        t.style = "Table Grid"

    for i, header in enumerate(table.headers):
        cell = t.rows[0].cells[i]
        cell.text = str(header)
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True

    for r, row in enumerate(table.rows, start=1):
        for c, value in enumerate(row[: len(table.headers)]):
            t.rows[r].cells[c].text = str(value)

    doc.add_paragraph()  # spacing after the table


def build_docx(plan: Plan, sections: list[Section]) -> bytes:
    doc = Document()

    # Title block
    title = doc.add_heading(plan.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"Document type: {plan.doc_type.replace('_', ' ').title()}")
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Assumptions
    if plan.assumptions:
        doc.add_heading("Assumptions", level=1)
        for assumption in plan.assumptions:
            doc.add_paragraph(assumption, style="List Bullet")

    # Sections
    for section in sections:
        doc.add_heading(section.heading, level=1)
        for paragraph in section.content.split("\n"):
            paragraph = paragraph.strip()
            if paragraph:
                doc.add_paragraph(paragraph)
        if section.table:
            _add_table(doc, section.table)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
